import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml

# Initialize presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# DESIGN SYSTEM COLOR PALETTE
COLOR_BG = RGBColor(248, 249, 250)         # Soft clean off-white background
COLOR_PRIMARY = RGBColor(27, 94, 32)        # Rich Forest Green
COLOR_SECONDARY = RGBColor(13, 71, 161)     # Deep Tech Blue
COLOR_TEXT_MAIN = RGBColor(33, 33, 33)      # Dark Charcoal Text
COLOR_TEXT_MUTED = RGBColor(117, 117, 117)  # Medium Muted Gray
COLOR_WHITE = RGBColor(255, 255, 255)       # Card Background
COLOR_CARD_GREEN = RGBColor(240, 248, 241)  # Soft Mint Green Card
COLOR_CARD_BLUE = RGBColor(240, 246, 252)   # Soft Sky Blue Card
COLOR_CARD_BORDER = RGBColor(224, 224, 224) # Subtle Gray Border
COLOR_ACCENT_GREEN = RGBColor(200, 230, 201) # Active Green Tag
COLOR_ACCENT_BLUE = RGBColor(187, 222, 251)  # Active Blue Tag

# Image Paths
TITLE_IMAGE_PATH = r"C:\Users\Mushfiqur Rahman\.gemini\antigravity\brain\8401e8a6-931d-4799-bf34-49bcd9038a68\title_slide_visual_1780668788451.png"
THANK_YOU_IMAGE_PATH = r"C:\Users\Mushfiqur Rahman\.gemini\antigravity\brain\8401e8a6-931d-4799-bf34-49bcd9038a68\thank_you_visual_1780668816512.png"

# Helpers
def add_slide_transition_fade(slide):
    """Applies a clean Fade transition to the slide via XML injection."""
    sld = slide._element
    transition_xml = (
        r'<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        r'<p:fade/>'
        r'</p:transition>'
    )
    sld.append(parse_xml(transition_xml))

def apply_background(slide, color=COLOR_BG):
    """Fills the slide background with a solid color."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_header_separator(slide):
    """Draws a thin modern separator line under the header."""
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.25), Inches(11.73), Inches(0.015))
    line.fill.solid()
    line.fill.fore_color.rgb = COLOR_CARD_BORDER
    line.line.fill.background()

def add_slide_header(slide, title_text, category_text="PROJECT PROPOSAL"):
    """Adds a standard structured header with a category tag, main title, and thin separator line."""
    # Category Tag
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.73), Inches(0.35))
    cat_tf = cat_box.text_frame
    cat_tf.word_wrap = True
    cat_tf.margin_left = cat_tf.margin_right = cat_tf.margin_top = cat_tf.margin_bottom = 0
    p_cat = cat_tf.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.name = "Segoe UI"
    p_cat.font.size = Pt(9.5)
    p_cat.font.bold = True
    p_cat.font.color.rgb = COLOR_SECONDARY

    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.73), Inches(0.65))
    title_tf = title_box.text_frame
    title_tf.word_wrap = True
    title_tf.margin_left = title_tf.margin_right = title_tf.margin_top = title_tf.margin_bottom = 0
    p_title = title_tf.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = "Segoe UI"
    p_title.font.size = Pt(25)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_PRIMARY

    add_header_separator(slide)

def add_rounded_card(slide, left, top, width, height, fill_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER, border_width=1):
    """Draws a clean, styled rounded rectangle card to container content."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = fill_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(border_width)
    else:
        card.line.fill.background()
    return card

def add_text_box(slide, left, top, width, height, text_paragraphs, word_wrap=True):
    """Utility to place rich text blocks within cards with complete style control."""
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = word_wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    
    for i, p_info in enumerate(text_paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = p_info.get("text", "")
        p.font.name = p_info.get("font_name", "Calibri")
        p.font.size = Pt(p_info.get("size", 12))
        p.font.bold = p_info.get("bold", False)
        p.font.italic = p_info.get("italic", False)
        p.font.color.rgb = p_info.get("color", COLOR_TEXT_MAIN)
        p.alignment = p_info.get("alignment", PP_ALIGN.LEFT)
        if "space_after" in p_info:
            p.space_after = Pt(p_info["space_after"])
        if "space_before" in p_info:
            p.space_before = Pt(p_info["space_before"])
        if "line_spacing" in p_info:
            p.line_spacing = p_info["line_spacing"]
    return tb

def add_page_number(slide, current_slide, total_slides=15):
    """Adds a clean slide number in the bottom right corner of body slides."""
    if current_slide == 1:
        return  # Skip title slide
    num_box = slide.shapes.add_textbox(Inches(11.5), Inches(6.9), Inches(1.0), Inches(0.3))
    num_tf = num_box.text_frame
    num_tf.word_wrap = True
    num_tf.margin_left = num_tf.margin_right = num_tf.margin_top = num_tf.margin_bottom = 0
    p = num_tf.paragraphs[0]
    p.text = f"{current_slide:02d} / {total_slides:02d}"
    p.font.name = "Segoe UI"
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.alignment = PP_ALIGN.RIGHT

# ==========================================
# SLIDE 1: TITLE SLIDE (Modern Split Design - Square Visual Center)
# ==========================================
slide1 = prs.slides.add_slide(prs.slide_layouts[6])
apply_background(slide1, COLOR_WHITE)
add_slide_transition_fade(slide1)

# Left Side Details
title_card = add_rounded_card(slide1, Inches(0.8), Inches(0.8), Inches(6.0), Inches(5.9), fill_color=COLOR_BG, border_color=None)
# Green vertical accent bar
accent_bar1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.8), Inches(0.12), Inches(5.9))
accent_bar1.fill.solid(); accent_bar1.fill.fore_color.rgb = COLOR_PRIMARY; accent_bar1.line.fill.background()

add_text_box(slide1, Inches(1.2), Inches(1.2), Inches(5.2), Inches(5.1), [
    {"text": "FINAL YEAR CSE PROJECT PROPOSAL", "font_name": "Segoe UI", "size": 11, "bold": True, "color": COLOR_SECONDARY, "space_after": 15},
    {"text": "IoT-Based Smart Soil\nAnalysis and AI-Powered\nCrop Recommendation\nSystem", "font_name": "Segoe UI", "size": 28, "bold": True, "color": COLOR_PRIMARY, "space_after": 25},
    {"text": "STUDENTS:", "font_name": "Segoe UI", "size": 10, "bold": True, "color": COLOR_TEXT_MUTED, "space_after": 4},
    {"text": "▪ Jeba Takia (210104)\n▪ MD Mushfiqur Rahman Mahim (210131)", "font_name": "Calibri", "size": 13, "bold": True, "color": COLOR_TEXT_MAIN, "space_after": 20},
    {"text": "INSTITUTION:", "font_name": "Segoe UI", "size": 10, "bold": True, "color": COLOR_TEXT_MUTED, "space_after": 4},
    {"text": "Department of Computer Science and Engineering\nDaffodil Institute of IT\nSession: 2020-21", "font_name": "Calibri", "size": 12, "color": COLOR_TEXT_MAIN}
])

# Right Side Image (Perfect 1:1 Aspect Ratio centered vertically inside the panel)
if os.path.exists(TITLE_IMAGE_PATH):
    img_card = add_rounded_card(slide1, Inches(7.1), Inches(0.8), Inches(5.4), Inches(5.9), fill_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER, border_width=2)
    # 5.2" x 5.2" image center coordinates: left = 7.1 + (5.4 - 5.2)/2 = 7.2; top = 0.8 + (5.9 - 5.2)/2 = 1.15
    slide1.shapes.add_picture(TITLE_IMAGE_PATH, Inches(7.2), Inches(1.15), Inches(5.2), Inches(5.2))

add_page_number(slide1, 1)

# ==========================================
# SLIDE 2: INTRODUCTION (Split Theme Analysis)
# ==========================================
slide2 = prs.slides.add_slide(prs.slide_layouts[6])
apply_background(slide2)
add_slide_transition_fade(slide2)
add_slide_header(slide2, "Introduction to Smart Agriculture")

# Column 1 (Smart Ag Card - Green Accent)
add_rounded_card(slide2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.0), fill_color=COLOR_CARD_GREEN)
add_rounded_card(slide2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(0.12), fill_color=COLOR_PRIMARY, border_color=None)
add_text_box(slide2, Inches(1.2), Inches(2.0), Inches(4.8), Inches(4.2), [
    {"text": "The Paradigm Shift", "font_name": "Segoe UI", "size": 18, "bold": True, "color": COLOR_PRIMARY, "space_after": 14},
    {"text": "Smart agriculture integrates digital technologies, real-time sensing, and analytical modeling to automate and optimize farming practices. As global food demand rises, farming must evolve from a reactive industry to a predictive, data-driven science.", "font_name": "Calibri", "size": 13, "color": COLOR_TEXT_MAIN, "space_after": 14},
    {"text": "Key Technology Pillars:", "font_name": "Segoe UI", "size": 12, "bold": True, "color": COLOR_SECONDARY, "space_after": 6},
    {"text": "▪ IoT Sensors: Continuous field telemetry (pH, Moisture, Temp).\n▪ Machine Learning: Smart predictive algorithms for cropping.\n▪ Cloud Integration: Centralized telemetry archives and API access.\n▪ Automation: Precision resource distribution (smart irrigation).", "font_name": "Calibri", "size": 12, "color": COLOR_TEXT_MAIN}
])

# Column 2 (Traditional Farming Card - Blue Accent)
add_rounded_card(slide2, Inches(6.93), Inches(1.6), Inches(5.6), Inches(5.0), fill_color=COLOR_CARD_BLUE)
add_rounded_card(slide2, Inches(6.93), Inches(1.6), Inches(5.6), Inches(0.12), fill_color=COLOR_SECONDARY, border_color=None)
add_text_box(slide2, Inches(7.33), Inches(2.0), Inches(4.8), Inches(4.2), [
    {"text": "Traditional Farming Bottlenecks", "font_name": "Segoe UI", "size": 18, "bold": True, "color": COLOR_SECONDARY, "space_after": 14},
    {"text": "Conventional farming practices in developing regions heavily depend on historical guesswork, visual cues, and generic seasonal patterns, leading to sub-optimal resource usage and crop failure.", "font_name": "Calibri", "size": 13, "color": COLOR_TEXT_MAIN, "space_after": 14},
    {"text": "Critical Challenges Faced:", "font_name": "Segoe UI", "size": 12, "bold": True, "color": COLOR_PRIMARY, "space_after": 6},
    {"text": "▪ Soil Nutrient Ignorance: Farming without knowing pH or moisture levels.\n▪ Resource Wastage: Over-watering or excessive use of chemical fertilizers.\n▪ Economic Risk: High crop failure rates due to mismatched soil compatibility.\n▪ Unpredictable Climate: Lack of adaptivity to shifting local weather.", "font_name": "Calibri", "size": 12, "color": COLOR_TEXT_MAIN}
])

add_page_number(slide2, 2)

# ==========================================
# SLIDE 3: PROBLEM STATEMENT (Alternating 3-Column Grid)
# ==========================================
slide3 = prs.slides.add_slide(prs.slide_layouts[6])
apply_background(slide3)
add_slide_transition_fade(slide3)
add_slide_header(slide3, "Problem Statement")

card_w = Inches(3.6)
card_h = Inches(4.8)
card_y = Inches(1.8)

# Problem 1 (Green Card)
add_rounded_card(slide3, Inches(0.8), card_y, card_w, card_h, fill_color=COLOR_CARD_GREEN)
add_rounded_card(slide3, Inches(0.8), card_y, card_w, Inches(0.12), fill_color=COLOR_PRIMARY, border_color=None)
add_text_box(slide3, Inches(1.1), Inches(2.2), Inches(3.0), Inches(4.0), [
    {"text": "01", "font_name": "Segoe UI", "size": 28, "bold": True, "color": COLOR_PRIMARY, "space_after": 10},
    {"text": "Guesswork Cultivation", "font_name": "Segoe UI", "size": 16, "bold": True, "color": COLOR_TEXT_MAIN, "space_after": 12},
    {"text": "Farmers often grow traditional crops year after year without validating if the soil has degraded. This lack of crop-to-soil compatibility testing directly leads to low yield rates and agricultural economic strain.", "font_name": "Calibri", "size": 12, "color": COLOR_TEXT_MAIN}
])

# Problem 2 (Blue Card)
add_rounded_card(slide3, Inches(4.86), card_y, card_w, card_h, fill_color=COLOR_CARD_BLUE)
add_rounded_card(slide3, Inches(4.86), card_y, card_w, Inches(0.12), fill_color=COLOR_SECONDARY, border_color=None)
add_text_box(slide3, Inches(5.16), Inches(2.2), Inches(3.0), Inches(4.0), [
    {"text": "02", "font_name": "Segoe UI", "size": 28, "bold": True, "color": COLOR_SECONDARY, "space_after": 10},
    {"text": "No Soil Telemetry Info", "font_name": "Segoe UI", "size": 16, "bold": True, "color": COLOR_TEXT_MAIN, "space_after": 12},
    {"text": "Soil properties fluctuate dynamically due to watering, fertilization, and climate. Lab-based soil testing takes weeks and is expensive, leaving farmers with no access to real-time, actionable chemical and moisture telemetry.", "font_name": "Calibri", "size": 12, "color": COLOR_TEXT_MAIN}
])

# Problem 3 (Green Card)
add_rounded_card(slide3, Inches(8.92), card_y, card_w, card_h, fill_color=COLOR_CARD_GREEN)
add_rounded_card(slide3, Inches(8.92), card_y, card_w, Inches(0.12), fill_color=COLOR_PRIMARY, border_color=None)
add_text_box(slide3, Inches(9.22), Inches(2.2), Inches(3.0), Inches(4.0), [
    {"text": "03", "font_name": "Segoe UI", "size": 28, "bold": True, "color": COLOR_PRIMARY, "space_after": 10},
    {"text": "Technology Adoption Gap", "font_name": "Segoe UI", "size": 16, "bold": True, "color": COLOR_TEXT_MAIN, "space_after": 12},
    {"text": "While advanced agricultural sensors and AI technologies exist, they are isolated from local farmers. There is a critical shortage of low-cost, integrated hardware and simple bilingual mobile-friendly dashboard systems.", "font_name": "Calibri", "size": 12, "color": COLOR_TEXT_MAIN}
])

add_page_number(slide3, 3)

# ==========================================
# SLIDE 4: PROJECT OBJECTIVES (Checkerboard 4-Quadrant Grid)
# ==========================================
slide4 = prs.slides.add_slide(prs.slide_layouts[6])
apply_background(slide4)
add_slide_transition_fade(slide4)
add_slide_header(slide4, "Project Objectives")

q_w = Inches(5.6)
q_h = Inches(2.3)

# Q1 (Top-Left - Green)
add_rounded_card(slide4, Inches(0.8), Inches(1.7), q_w, q_h, fill_color=COLOR_CARD_GREEN)
add_rounded_card(slide4, Inches(0.8), Inches(1.7), Inches(0.12), q_h, fill_color=COLOR_PRIMARY, border_color=None)
add_text_box(slide4, Inches(1.2), Inches(1.9), Inches(4.8), Inches(1.8), [
    {"text": "01. Real-Time Soil Monitoring", "font_name": "Segoe UI", "size": 16, "bold": True, "color": COLOR_PRIMARY, "space_after": 6},
    {"text": "Deploy low-cost, high-precision IoT sensors to continuously capture critical soil parameters: pH, moisture content, and ambient temperature/humidity.", "font_name": "Calibri", "size": 12, "color": COLOR_TEXT_MAIN}
])

# Q2 (Top-Right - Blue)
add_rounded_card(slide4, Inches(6.93), Inches(1.7), q_w, q_h, fill_color=COLOR_CARD_BLUE)
add_rounded_card(slide4, Inches(6.93), Inches(1.7), Inches(0.12), q_h, fill_color=COLOR_SECONDARY, border_color=None)
add_text_box(slide4, Inches(7.33), Inches(1.9), Inches(4.8), Inches(1.8), [
    {"text": "02. Cloud Database & Telemetry Logs", "font_name": "Segoe UI", "size": 16, "bold": True, "color": COLOR_SECONDARY, "space_after": 6},
    {"text": "Establish a secure, lightweight cloud database to log all real-time field readings, forming a historical timeline for analysis and system auditing.", "font_name": "Calibri", "size": 12, "color": COLOR_TEXT_MAIN}
])

# Q3 (Bottom-Left - Blue)
add_rounded_card(slide4, Inches(0.8), Inches(4.3), q_w, q_h, fill_color=COLOR_CARD_BLUE)
add_rounded_card(slide4, Inches(0.8), Inches(4.3), Inches(0.12), q_h, fill_color=COLOR_SECONDARY, border_color=None)
add_text_box(slide4, Inches(1.2), Inches(4.5), Inches(4.8), Inches(1.8), [
    {"text": "03. AI-Powered Crop Recommendation", "font_name": "Segoe UI", "size": 16, "bold": True, "color": COLOR_SECONDARY, "space_after": 6},
    {"text": "Design and train an ML classifier (Random Forest/SVM) that processes real-time sensor parameters to recommend the most suitable crop and required fertilizers.", "font_name": "Calibri", "size": 12, "color": COLOR_TEXT_MAIN}
])

# Q4 (Bottom-Right - Green)
add_rounded_card(slide4, Inches(6.93), Inches(4.3), q_w, q_h, fill_color=COLOR_CARD_GREEN)
add_rounded_card(slide4, Inches(6.93), Inches(4.3), Inches(0.12), q_h, fill_color=COLOR_PRIMARY, border_color=None)
add_text_box(slide4, Inches(7.33), Inches(4.5), Inches(4.8), Inches(1.8), [
    {"text": "04. Web Dashboard Visualization", "font_name": "Segoe UI", "size": 16, "bold": True, "color": COLOR_PRIMARY, "space_after": 6},
    {"text": "Develop an accessible web interface with real-time graphs, simple health scores, bilingual translation, and automated SMS alerts for soil degradation.", "font_name": "Calibri", "size": 12, "color": COLOR_TEXT_MAIN}
])

add_page_number(slide4, 4)

# ==========================================
# SLIDE 5: PROPOSED SYSTEM OVERVIEW (Architecture Flowchart)
# ==========================================
slide5 = prs.slides.add_slide(prs.slide_layouts[6])
apply_background(slide5)
add_slide_transition_fade(slide5)
add_slide_header(slide5, "Proposed System Architecture")

# Explanation at the top
add_text_box(slide5, Inches(0.8), Inches(1.4), Inches(11.7), Inches(0.5), [
    {"text": "The data starts at the sensors, is processed by ESP32, uploaded to the cloud database, analyzed by AI, and served to the farmer dashboard.", "font_name": "Calibri", "size": 12, "italic": True, "color": COLOR_TEXT_MUTED}
])

# Let's draw 5 block shapes horizontally with arrows in between.
box_w = Inches(1.8)
box_h = Inches(2.2)
box_y = Inches(2.5)
arrow_w = Inches(0.4)
arrow_h = Inches(0.3)
arrow_y = box_y + Inches(0.95)

# Shapes definition: (Left, Title, Color, Content)
blocks = [
    (Inches(0.8), "01. SENSORS", COLOR_PRIMARY, "pH Sensor\nSoil Moisture\nDHT22 Temp"),
    (Inches(3.2), "02. GATEWAY", COLOR_SECONDARY, "ESP32 MCU\nAnalog ADC\nWi-Fi Upload"),
    (Inches(5.6), "03. CLOUD DB", COLOR_PRIMARY, "MySQL / Postgre\nSecure API\nJSON Logging"),
    (Inches(8.0), "04. AI ENGINE", COLOR_SECONDARY, "Python ML API\nRandom Forest\nInference Model"),
    (Inches(10.4), "05. DASHBOARD", COLOR_PRIMARY, "HTML/CSS/JS\nReal-time Alerts\nVisual Graphs")
]

for idx, (bx, btitle, bcolor, bcontent) in enumerate(blocks):
    # Main container
    add_rounded_card(slide5, bx, box_y, box_w, box_h, fill_color=COLOR_WHITE, border_color=bcolor, border_width=2)
    # Header tag in container
    tag = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, bx, box_y, box_w, Inches(0.4))
    tag.fill.solid()
    tag.fill.fore_color.rgb = bcolor
    tag.line.fill.background()
    
    # Title Text inside header tag
    tf_tag = tag.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_right = tf_tag.margin_top = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = btitle
    p_tag.font.name = "Segoe UI"
    p_tag.font.size = Pt(9.5)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_WHITE
    p_tag.alignment = PP_ALIGN.CENTER
    
    # Content inside container
    add_text_box(slide5, bx + Inches(0.1), box_y + Inches(0.6), box_w - Inches(0.2), box_h - Inches(0.8), [
        {"text": bcontent, "font_name": "Calibri", "size": 11.5, "alignment": PP_ALIGN.CENTER, "color": COLOR_TEXT_MAIN, "line_spacing": 1.2}
    ])
    
    # Arrow (except for last box)
    if idx < 4:
        arrow_x = bx + box_w + Inches(0.1)
        arrow = slide5.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, arrow_x, arrow_y, arrow_w, arrow_h)
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = COLOR_TEXT_MUTED
        arrow.line.fill.background()

# Add a bottom summary box
summary_box = add_rounded_card(slide5, Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.3), fill_color=COLOR_CARD_BLUE, border_color=None)
add_text_box(slide5, Inches(1.2), Inches(5.5), Inches(11.0), Inches(0.9), [
    {"text": "Data Flow Summary:", "font_name": "Segoe UI", "size": 12, "bold": True, "color": COLOR_SECONDARY, "space_after": 4},
    {"text": "1. Soil values are captured analogously. 2. ESP32 digitizes and transmits via Wi-Fi (HTTP POST/MQTT). 3. Cloud server stores data points in a secure DB. 4. Flask API triggers inference. 5. Recommendations and real-time parameters update on the dashboard.", "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MAIN}
])

add_page_number(slide5, 5)

# ==========================================
# SLIDE 6: HARDWARE REQUIREMENTS (2x2 Grid + Details)
# ==========================================
slide6 = prs.slides.add_slide(prs.slide_layouts[6])
apply_background(slide6)
add_slide_transition_fade(slide6)
add_slide_header(slide6, "Hardware Architecture & Sensors")

grid_w = Inches(5.6)
grid_h = Inches(2.3)

# Row 1 Left - ESP32
add_rounded_card(slide6, Inches(0.8), Inches(1.6), grid_w, grid_h, fill_color=COLOR_CARD_GREEN)
add_rounded_card(slide6, Inches(0.8), Inches(1.6), Inches(0.12), grid_h, fill_color=COLOR_PRIMARY, border_color=None)
add_text_box(slide6, Inches(1.2), Inches(1.8), Inches(5.0), Inches(2.0), [
    {"text": "ESP32 NodeMCU Microcontroller", "font_name": "Segoe UI", "size": 15, "bold": True, "color": COLOR_PRIMARY, "space_after": 4},
    {"text": "▪ Core: Dual-core Tensilica LX6 32-bit MCU with integrated Wi-Fi and Bluetooth.\n▪ Pinouts: Multi-channel 12-bit ADC for sensor calibration and digital GPIOs.\n▪ Purpose: Acts as the central node for sensor calibration, signal digitizing, and HTTP data transmission.", "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MAIN}
])

# Row 1 Right - Soil Moisture
add_rounded_card(slide6, Inches(6.93), Inches(1.6), grid_w, grid_h, fill_color=COLOR_CARD_BLUE)
add_rounded_card(slide6, Inches(6.93), Inches(1.6), Inches(0.12), grid_h, fill_color=COLOR_SECONDARY, border_color=None)
add_text_box(slide6, Inches(7.33), Inches(1.8), Inches(5.0), Inches(2.0), [
    {"text": "Capacitive Soil Moisture Sensor (v1.2)", "font_name": "Segoe UI", "size": 15, "bold": True, "color": COLOR_SECONDARY, "space_after": 4},
    {"text": "▪ Tech: Capacitive measurements prevent probe oxidation and corrosion.\n▪ Principle: Measures dielectric permittivity of surrounding soil to estimate water content percentage.\n▪ Purpose: Signals irrigation frequency requirements.", "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MAIN}
])

# Row 2 Left - pH Sensor
add_rounded_card(slide6, Inches(0.8), Inches(4.2), grid_w, grid_h, fill_color=COLOR_CARD_BLUE)
add_rounded_card(slide6, Inches(0.8), Inches(4.2), Inches(0.12), grid_h, fill_color=COLOR_SECONDARY, border_color=None)
add_text_box(slide6, Inches(1.2), Inches(4.4), Inches(5.0), Inches(2.0), [
    {"text": "Soil pH Sensor Probe", "font_name": "Segoe UI", "size": 15, "bold": True, "color": COLOR_SECONDARY, "space_after": 4},
    {"text": "▪ Tech: High-precision glass electrode analog probe with signal amplifier.\n▪ Output: 0-5V analog signal proportional to hydrogen-ion activity (pH scale 0-14).\n▪ Purpose: Key driver for crop suitability and fertilizer recommendation.", "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MAIN}
])

# Row 2 Right - Temp & Humidity
add_rounded_card(slide6, Inches(6.93), Inches(4.2), grid_w, grid_h, fill_color=COLOR_CARD_GREEN)
add_rounded_card(slide6, Inches(6.93), Inches(4.2), Inches(0.12), grid_h, fill_color=COLOR_PRIMARY, border_color=None)
add_text_box(slide6, Inches(7.33), Inches(4.4), Inches(5.0), Inches(2.0), [
    {"text": "DHT22 Temperature & Humidity Sensor", "font_name": "Segoe UI", "size": 15, "bold": True, "color": COLOR_PRIMARY, "space_after": 4},
    {"text": "▪ Specs: Temperature range -40 to 80°C (+-0.5°C); Humidity range 0-100% (+-2% RH).\n▪ Digital: Single-bus interface delivers digitized environmental parameters.\n▪ Purpose: Analyzes microclimate to determine transpiration index.", "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MAIN}
])

add_page_number(slide6, 6)

# ==========================================
# SLIDE 7: SOFTWARE REQUIREMENTS (Split Grid)
# ==========================================
slide7 = prs.slides.add_slide(prs.slide_layouts[6])
apply_background(slide7)
add_slide_transition_fade(slide7)
add_slide_header(slide7, "Software Environment & Core Stack")

col_w = Inches(5.6)
col_h = Inches(5.0)

# Left Side (Development & Backend Stack)
add_rounded_card(slide7, Inches(0.8), Inches(1.7), col_w, col_h, fill_color=COLOR_CARD_GREEN)
add_rounded_card(slide7, Inches(0.8), Inches(1.7), col_w, Inches(0.12), fill_color=COLOR_PRIMARY, border_color=None)
add_text_box(slide7, Inches(1.2), Inches(2.1), Inches(4.8), Inches(4.2), [
    {"text": "Firmware & Backend Pipeline", "font_name": "Segoe UI", "size": 18, "bold": True, "color": COLOR_PRIMARY, "space_after": 15},
    {"text": "▪ Arduino IDE (C/C++): Program ESP32 for analog calibration, digital filtering, Wi-Fi reconnection logic, and REST API transmission.\n\n▪ Python Flask / Node.js API: Lightweight backend server. Handles incoming JSON sensor payloads, routes requests, and exposes REST endpoints for the client dashboard.\n\n▪ MySQL / PostgreSQL Database: Relational structure with tables storing client profile details, ESP32 device keys, logged soil data points, and recommendations log.", "font_name": "Calibri", "size": 12, "color": COLOR_TEXT_MAIN}
])

# Right Side (Frontend Stack)
add_rounded_card(slide7, Inches(6.93), Inches(1.7), col_w, col_h, fill_color=COLOR_CARD_BLUE)
add_rounded_card(slide7, Inches(6.93), Inches(1.7), col_w, Inches(0.12), fill_color=COLOR_SECONDARY, border_color=None)
add_text_box(slide7, Inches(7.33), Inches(2.1), Inches(4.8), Inches(4.2), [
    {"text": "Visualization & ML Pipeline", "font_name": "Segoe UI", "size": 18, "bold": True, "color": COLOR_SECONDARY, "space_after": 15},
    {"text": "▪ HTML5, CSS3, JavaScript (ES6+): Responsive front-end built using CSS grid, flexbox, and Chart.js for real-time visualization graphs.\n\n▪ Machine Learning Suite (Scikit-Learn, Pandas, Numpy): Used to clean agricultural soil datasets, perform feature scaling, train classifiers, and serialize the model for Flask integration.\n\n▪ Twilio / SMS API Gateway: Integrates custom notification loops to automatically push critical SMS alerts to registered farmers during extreme conditions.", "font_name": "Calibri", "size": 12, "color": COLOR_TEXT_MAIN}
])

add_page_number(slide7, 7)

# ==========================================
# SLIDE 8: WORKFLOW DIAGRAM (Step-by-Step Flow - Centered Coordinates)
# ==========================================
slide8 = prs.slides.add_slide(prs.slide_layouts[6])
apply_background(slide8)
add_slide_transition_fade(slide8)
add_slide_header(slide8, "System Workflow & Process Pipeline")

# We will layout a 3x2 grid of workflow blocks.
wf_w = Inches(3.6)
wf_h = Inches(2.2)

# Centered X coordinates for 3 columns:
col1_x = Inches(0.8)
col2_x = Inches(4.86)
col3_x = Inches(8.92)

# Row 1
# Step 1
add_rounded_card(slide8, col1_x, Inches(1.8), wf_w, wf_h, fill_color=COLOR_CARD_GREEN)
add_text_box(slide8, col1_x + Inches(0.3), Inches(2.0), wf_w - Inches(0.6), Inches(1.8), [
    {"text": "Step 1: Data Sensing", "font_name": "Segoe UI", "size": 14, "bold": True, "color": COLOR_PRIMARY, "space_after": 4},
    {"text": "Sensors sample pH, moisture, and temperature ambient levels continuously. Analog values are filtered to remove background noise.", "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MAIN}
])

# Step 2
add_rounded_card(slide8, col2_x, Inches(1.8), wf_w, wf_h, fill_color=COLOR_CARD_BLUE)
add_text_box(slide8, col2_x + Inches(0.3), Inches(2.0), wf_w - Inches(0.6), Inches(1.8), [
    {"text": "Step 2: Transmission", "font_name": "Segoe UI", "size": 14, "bold": True, "color": COLOR_SECONDARY, "space_after": 4},
    {"text": "ESP32 bundles sensor readings into a JSON payload and transmits it securely via HTTP POST request to the cloud endpoint.", "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MAIN}
])

# Step 3
add_rounded_card(slide8, col3_x, Inches(1.8), wf_w, wf_h, fill_color=COLOR_CARD_GREEN)
add_text_box(slide8, col3_x + Inches(0.3), Inches(2.0), wf_w - Inches(0.6), Inches(1.8), [
    {"text": "Step 3: Cloud Logging", "font_name": "Segoe UI", "size": 14, "bold": True, "color": COLOR_PRIMARY, "space_after": 4},
    {"text": "The Flask server parses the request, validates the device token, and logs the telemetry with timestamps into the SQL database.", "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MAIN}
])

# Row 2
# Step 4
add_rounded_card(slide8, col1_x, Inches(4.4), wf_w, wf_h, fill_color=COLOR_CARD_BLUE)
add_text_box(slide8, col1_x + Inches(0.3), Inches(4.6), wf_w - Inches(0.6), Inches(1.8), [
    {"text": "Step 4: AI Model Inference", "font_name": "Segoe UI", "size": 14, "bold": True, "color": COLOR_SECONDARY, "space_after": 4},
    {"text": "The server triggers the ML inference engine. The pre-trained classifier matches soil parameters against crop compatibility arrays.", "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MAIN}
])

# Step 5
add_rounded_card(slide8, col2_x, Inches(4.4), wf_w, wf_h, fill_color=COLOR_CARD_GREEN)
add_text_box(slide8, col2_x + Inches(0.3), Inches(4.6), wf_w - Inches(0.6), Inches(1.8), [
    {"text": "Step 5: Recommendation", "font_name": "Segoe UI", "size": 14, "bold": True, "color": COLOR_PRIMARY, "space_after": 4},
    {"text": "AI outputs the optimal crop type along with specific fertilizer ratios. Warnings are triggered if severe anomalies are detected.", "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MAIN}
])

# Step 6
add_rounded_card(slide8, col3_x, Inches(4.4), wf_w, wf_h, fill_color=COLOR_CARD_BLUE)
add_text_box(slide8, col3_x + Inches(0.3), Inches(4.6), wf_w - Inches(0.6), Inches(1.8), [
    {"text": "Step 6: Dashboard Update", "font_name": "Segoe UI", "size": 14, "bold": True, "color": COLOR_SECONDARY, "space_after": 4},
    {"text": "The Web Dashboard pulls fresh logs via AJAX/WebSockets. Charts re-render, and crop advisory reports update instantly.", "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MAIN}
])

# Draw connector arrows (horizontal)
# Row 1 Arrow 1 (Between 4.4 and 4.86, center is 4.63)
arrow1 = slide8.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.5), Inches(2.75), Inches(0.26), Inches(0.2))
arrow1.fill.solid(); arrow1.fill.fore_color.rgb = COLOR_TEXT_MUTED; arrow1.line.fill.background()
# Row 1 Arrow 2 (Between 8.46 and 8.92, center is 8.69)
arrow2 = slide8.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.56), Inches(2.75), Inches(0.26), Inches(0.2))
arrow2.fill.solid(); arrow2.fill.fore_color.rgb = COLOR_TEXT_MUTED; arrow2.line.fill.background()

# Arrow down from Step 3 to Step 6 (Step 3/6 center: 8.92 + 1.8 = 10.72)
arrow_down = slide8.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(10.62), Inches(4.1), Inches(0.2), Inches(0.25))
arrow_down.fill.solid(); arrow_down.fill.fore_color.rgb = COLOR_TEXT_MUTED; arrow_down.line.fill.background()

# Row 2 Arrow 1 (Leftwards from Step 6 to Step 5)
arrow3 = slide8.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, Inches(8.56), Inches(5.35), Inches(0.26), Inches(0.2))
arrow3.fill.solid(); arrow3.fill.fore_color.rgb = COLOR_TEXT_MUTED; arrow3.line.fill.background()
# Row 2 Arrow 2 (Leftwards from Step 5 to Step 4)
arrow4 = slide8.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, Inches(4.5), Inches(5.35), Inches(0.26), Inches(0.2))
arrow4.fill.solid(); arrow4.fill.fore_color.rgb = COLOR_TEXT_MUTED; arrow4.line.fill.background()

add_page_number(slide8, 8)

# ==========================================
# SLIDE 9: KEY FEATURES (3x2 Dashboard Grid)
# ==========================================
slide9 = prs.slides.add_slide(prs.slide_layouts[6])
apply_background(slide9)
add_slide_transition_fade(slide9)
add_slide_header(slide9, "Key Features & Platform Capabilities")

kf_w = Inches(3.6)
kf_h = Inches(2.3)

# Grid coordinate alignments
col1_x = Inches(0.8)
col2_x = Inches(4.86)
col3_x = Inches(8.92)

features = [
    (col1_x, Inches(1.7), "Real-time Telemetry", "Provides continuous logging of essential soil chemistry and ambient readings with graphical curves.", COLOR_CARD_GREEN, COLOR_PRIMARY),
    (col2_x, Inches(1.7), "AI Recommendation", "Utilizes classification algorithms matching multi-parameter inputs against crops with high precision.", COLOR_CARD_BLUE, COLOR_SECONDARY),
    (col3_x, Inches(1.7), "Fertilizer Optimizer", "Calculates required NPK adjustments based on the difference between present soil and crop ideals.", COLOR_CARD_GREEN, COLOR_PRIMARY),
    (col1_x, Inches(4.3), "Irrigation Advisory", "Automated alert recommendations for watering based on current moisture content and air humidity profiles.", COLOR_CARD_BLUE, COLOR_SECONDARY),
    (col2_x, Inches(4.3), "Soil Health Scoring", "Aggregates pH variance, water logging, and moisture trends into a unified, clean Health Score (1-100).", COLOR_CARD_GREEN, COLOR_PRIMARY),
    (col3_x, Inches(4.3), "Bilingual Dashboard", "Features toggleable English & Bangla views and mobile layout optimization for rural farmers.", COLOR_CARD_BLUE, COLOR_SECONDARY),
]

for fx, fy, ftitle, fdesc, fbg, fcol in features:
    add_rounded_card(slide9, fx, fy, kf_w, kf_h, fill_color=fbg)
    add_rounded_card(slide9, fx, fy, Inches(0.12), kf_h, fill_color=fcol, border_color=None)
    add_text_box(slide9, fx + Inches(0.3), fy + Inches(0.3), kf_w - Inches(0.4), kf_h - Inches(0.6), [
        {"text": ftitle, "font_name": "Segoe UI", "size": 15, "bold": True, "color": COLOR_TEXT_MAIN, "space_after": 6},
        {"text": fdesc, "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MUTED}
    ])

add_page_number(slide9, 9)

# ==========================================
# SLIDE 10: AI RECOMMENDATION MODULE (Flow & Model)
# ==========================================
slide10 = prs.slides.add_slide(prs.slide_layouts[6])
apply_background(slide10)
add_slide_transition_fade(slide10)
add_slide_header(slide10, "AI Crop Recommendation Engine")

# Left Column (Concept & Architecture)
add_rounded_card(slide10, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.0), fill_color=COLOR_CARD_GREEN)
add_rounded_card(slide10, Inches(0.8), Inches(1.7), Inches(5.6), Inches(0.12), fill_color=COLOR_PRIMARY, border_color=None)
add_text_box(slide10, Inches(1.2), Inches(2.1), Inches(4.8), Inches(4.3), [
    {"text": "Decision Logic & Classifiers", "font_name": "Segoe UI", "size": 18, "bold": True, "color": COLOR_PRIMARY, "space_after": 12},
    {"text": "The AI module uses supervised learning to map soil parameters to the target crop. Machine learning models analyze multi-dimensional bounds to generate recommendations.", "font_name": "Calibri", "size": 12.5, "color": COLOR_TEXT_MAIN, "space_after": 14},
    {"text": "Core ML Classifiers Evaluated:", "font_name": "Segoe UI", "size": 12, "bold": True, "color": COLOR_SECONDARY, "space_after": 6},
    {"text": "▪ Random Forest Classifier: High accuracy, robust against outliers, and works efficiently with structured soil data.\n▪ Support Vector Machine (SVM): Strong margin separation bounds.\n▪ Gaussian Naive Bayes: Fast inference with probabilistic outputs.\n▪ Decision Tree: Transparent routing logic matching local farming rules.", "font_name": "Calibri", "size": 11.5, "color": COLOR_TEXT_MAIN}
])

# Right Column (Decision Flowchart Diagram)
add_rounded_card(slide10, Inches(6.93), Inches(1.7), Inches(5.6), Inches(5.0), fill_color=COLOR_CARD_BLUE)
add_rounded_card(slide10, Inches(6.93), Inches(1.7), Inches(5.6), Inches(0.12), fill_color=COLOR_SECONDARY, border_color=None)

# Flowchart inside right card (X coordinate centered on right card center: 6.93 + 2.8 = 9.73)
box_x = Inches(7.33)
box_w = Inches(4.8)

# Step 1: Input parameters box
inp_box = add_rounded_card(slide10, box_x, Inches(2.2), box_w, Inches(0.9), fill_color=COLOR_ACCENT_GREEN, border_color=None)
add_text_box(slide10, box_x + Inches(0.2), Inches(2.3), box_w - Inches(0.4), Inches(0.7), [
    {"text": "INPUT PARAMETERS", "font_name": "Segoe UI", "size": 9.5, "bold": True, "color": COLOR_PRIMARY},
    {"text": "pH, Moisture %, NPK Ratios, Temperature, Humidity", "font_name": "Calibri", "size": 11, "bold": True, "color": COLOR_TEXT_MAIN}
])

arrow_d1 = slide10.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(9.63), Inches(3.2), Inches(0.2), Inches(0.3))
arrow_d1.fill.solid(); arrow_d1.fill.fore_color.rgb = COLOR_TEXT_MUTED; arrow_d1.line.fill.background()

# Step 2: Classifier Box
ml_box = add_rounded_card(slide10, box_x, Inches(3.6), box_w, Inches(0.9), fill_color=COLOR_ACCENT_BLUE, border_color=None)
add_text_box(slide10, box_x + Inches(0.2), Inches(3.7), box_w - Inches(0.4), Inches(0.7), [
    {"text": "RANDOM FOREST / SVM CLASSIFIER MODEL", "font_name": "Segoe UI", "size": 9.5, "bold": True, "color": COLOR_SECONDARY},
    {"text": "Calculates probabilistic score matrices for 22+ crop categories", "font_name": "Calibri", "size": 11, "bold": True, "color": COLOR_TEXT_MAIN}
])

arrow_d2 = slide10.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(9.63), Inches(4.6), Inches(0.2), Inches(0.3))
arrow_d2.fill.solid(); arrow_d2.fill.fore_color.rgb = COLOR_TEXT_MUTED; arrow_d2.line.fill.background()

# Step 3: Recommendation Outputs
out_box = add_rounded_card(slide10, box_x, Inches(5.0), box_w, Inches(1.2), fill_color=COLOR_WHITE, border_color=COLOR_CARD_BORDER, border_width=2)
# Green accent inside output
add_rounded_card(slide10, box_x, Inches(5.0), Inches(0.12), Inches(1.2), fill_color=COLOR_PRIMARY, border_color=None)
add_text_box(slide10, box_x + Inches(0.3), Inches(5.1), box_w - Inches(0.4), Inches(1.0), [
    {"text": "GENERATED SYSTEM OUTPUTS", "font_name": "Segoe UI", "size": 9.5, "bold": True, "color": COLOR_PRIMARY, "space_after": 2},
    {"text": "1. Recommended Crop: Rice, Maize, Lentils, Cotton, etc.\n2. Fertilizer Advisory: Specific NPK ratio adjustments.\n3. Irrigation Command: Active watering schedule.", "font_name": "Calibri", "size": 10.5, "color": COLOR_TEXT_MAIN}
])

add_page_number(slide10, 10)

# ==========================================
# SLIDE 11: BUSINESS PROSPECT (3 Strategic Pillars)
# ==========================================
slide11 = prs.slides.add_slide(prs.slide_layouts[6])
apply_background(slide11)
add_slide_transition_fade(slide11)
add_slide_header(slide11, "Business Model & Strategic Prospects")

bp_w = Inches(3.6)
bp_h = Inches(4.8)
bp_y = Inches(1.8)

# Alignments matching key columns
col1_x = Inches(0.8)
col2_x = Inches(4.86)
col3_x = Inches(8.92)

# Pillar 1: Smart Hardware Kit (Green Accent)
add_rounded_card(slide11, col1_x, bp_y, bp_w, bp_h, fill_color=COLOR_CARD_GREEN)
add_rounded_card(slide11, col1_x, bp_y, bp_w, Inches(0.12), fill_color=COLOR_PRIMARY, border_color=None)
add_text_box(slide11, col1_x + Inches(0.3), Inches(2.2), bp_w - Inches(0.6), Inches(4.0), [
    {"text": "PILLAR 1", "font_name": "Segoe UI", "size": 11, "bold": True, "color": COLOR_PRIMARY, "space_after": 8},
    {"text": "Smart Farming Kit", "font_name": "Segoe UI", "size": 16, "bold": True, "color": COLOR_TEXT_MAIN, "space_after": 14},
    {"text": "▪ One-Time Purchase: Sell localized, pre-calibrated ESP32 and sensor enclosures directly to commercial farms and local crop growers.\n\n▪ Plug-and-Play Setup: Users set up the device and connect to local Wi-Fi without needing technical support.", "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MAIN}
])

# Pillar 2: SaaS Dashboard Subscriptions (Blue Accent)
add_rounded_card(slide11, col2_x, bp_y, bp_w, bp_h, fill_color=COLOR_CARD_BLUE)
add_rounded_card(slide11, col2_x, bp_y, bp_w, Inches(0.12), fill_color=COLOR_SECONDARY, border_color=None)
add_text_box(slide11, col2_x + Inches(0.3), Inches(2.2), bp_w - Inches(0.6), Inches(4.0), [
    {"text": "PILLAR 2", "font_name": "Segoe UI", "size": 11, "bold": True, "color": COLOR_SECONDARY, "space_after": 8},
    {"text": "SaaS Subscription", "font_name": "Segoe UI", "size": 16, "bold": True, "color": COLOR_TEXT_MAIN, "space_after": 14},
    {"text": "▪ Tiered Subscription Model: Free basic level showing parameters. Premium level offering historical trend analysis, PDF soil health cards, and multi-device management.\n\n▪ Fertilizer API Access: Access for commercial fertilizer distributors looking to target specific regions.", "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MAIN}
])

# Pillar 3: B2G & NGO Partnerships (Green Accent)
add_rounded_card(slide11, col3_x, bp_y, bp_w, bp_h, fill_color=COLOR_CARD_GREEN)
add_rounded_card(slide11, col3_x, bp_y, bp_w, Inches(0.12), fill_color=COLOR_PRIMARY, border_color=None)
add_text_box(slide11, col3_x + Inches(0.3), Inches(2.2), bp_w - Inches(0.6), Inches(4.0), [
    {"text": "PILLAR 3", "font_name": "Segoe UI", "size": 11, "bold": True, "color": COLOR_PRIMARY, "space_after": 8},
    {"text": "NGO & Government Contracts", "font_name": "Segoe UI", "size": 16, "bold": True, "color": COLOR_TEXT_MAIN, "space_after": 14},
    {"text": "▪ Subsidized Distribution: Collaborate with agricultural ministries and microfinance NGOs to distribute kits to rural farming cooperatives.\n\n▪ Regional Data Dashboard: Package regional soil mapping data for policy makers planning agricultural support.", "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MAIN}
])

add_page_number(slide11, 11)

# ==========================================
# SLIDE 12: FUTURE ENHANCEMENTS (Horizontal Roadmap)
# ==========================================
slide12 = prs.slides.add_slide(prs.slide_layouts[6])
apply_background(slide12)
add_slide_transition_fade(slide12)
add_slide_header(slide12, "Future Roadmap & Enhancements")

# Horizontal Line for Roadmap
roadmap_line = slide12.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.6), Inches(11.73), Inches(0.06))
roadmap_line.fill.solid(); roadmap_line.fill.fore_color.rgb = COLOR_TEXT_MUTED; roadmap_line.line.fill.background()

# Let's place 4 roadmap points
p_w = Inches(2.6)
p_h = Inches(2.0)

# Centered Box X coordinates:
box1_x = Inches(0.8)
box2_x = Inches(3.84)
box3_x = Inches(6.88)
box4_x = Inches(9.92)

# Nodes centered coordinate offsets: box_x + (2.6 - 0.3)/2 = box_x + 1.15
positions = [
    (box1_x, Inches(1.4), Inches(1.95), "Phase 1: Weather API", "Incorporate live satellite weather forecasting (rainfall probability, frost warnings) into recommendation loop.", COLOR_PRIMARY, COLOR_CARD_GREEN),
    (box2_x, Inches(3.9), Inches(4.99), "Phase 2: Mobile App", "Develop a native mobile application using React Native to offer push alerts and offline telemetry caching.", COLOR_SECONDARY, COLOR_CARD_BLUE),
    (box3_x, Inches(1.4), Inches(8.03), "Phase 3: Smart Irrigation", "Integrate a solenoid valve circuit controller to enable automatic irrigation when moisture falls below thresholds.", COLOR_PRIMARY, COLOR_CARD_GREEN),
    (box4_x, Inches(3.9), Inches(11.07), "Phase 4: NPK Sensors", "Replace manual estimates with spectrophotometric optical sensors that measure soil NPK levels in real time.", COLOR_SECONDARY, COLOR_CARD_BLUE)
]

for idx, (px, py, node_x, ptitle, pdesc, node_color, card_bg) in enumerate(positions):
    # Circle node on line
    node = slide12.shapes.add_shape(MSO_SHAPE.OVAL, node_x, Inches(3.48), Inches(0.3), Inches(0.3))
    node.fill.solid(); node.fill.fore_color.rgb = node_color; node.line.fill.background()
    
    # Text container card
    add_rounded_card(slide12, px, py, p_w, p_h, fill_color=card_bg)
    add_rounded_card(slide12, px, py, p_w, Inches(0.08), fill_color=node_color, border_color=None)
    add_text_box(slide12, px + Inches(0.2), py + Inches(0.2), p_w - Inches(0.4), p_h - Inches(0.4), [
        {"text": ptitle, "font_name": "Segoe UI", "size": 12.5, "bold": True, "color": COLOR_TEXT_MAIN, "space_after": 6},
        {"text": pdesc, "font_name": "Calibri", "size": 10, "color": COLOR_TEXT_MAIN}
    ])

add_page_number(slide12, 12)

# ==========================================
# SLIDE 13: GANTT CHART / PROJECT TIMELINE (Custom Table)
# ==========================================
slide13 = prs.slides.add_slide(prs.slide_layouts[6])
apply_background(slide13)
add_slide_transition_fade(slide13)
add_slide_header(slide13, "Project Development Timeline")

# We will create a clean tabular Gantt chart
rows, cols = 7, 5
table_left = Inches(0.8)
table_top = Inches(1.8)
table_width = Inches(11.73)
table_height = Inches(4.5)

table_shape = slide13.shapes.add_table(rows, cols, table_left, table_top, table_width, table_height)
table = table_shape.table

# Column widths
table.columns[0].width = Inches(3.73)
for col_idx in range(1, 5):
    table.columns[col_idx].width = Inches(2.0)

# Headers
headers = ["Tasks / Phases", "Month 1 (Planning)", "Month 2 (HW Setup)", "Month 3 (AI & Backend)", "Month 4 (Dashboard & QA)"]
for col_idx, text in enumerate(headers):
    cell = table.cell(0, col_idx)
    cell.fill.solid(); cell.fill.fore_color.rgb = COLOR_PRIMARY
    # Add text
    tf = cell.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Inches(0.1)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Segoe UI"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE
    p.alignment = PP_ALIGN.CENTER if col_idx > 0 else PP_ALIGN.LEFT

# Row contents: (Task, Span start col, Span end col)
timeline_data = [
    ("Literature Review & Proposal", 1, 1),
    ("Hardware Calibration & Driver Config", 2, 2),
    ("AI Model Training & Dataset Cleaning", 2, 3),
    ("Cloud Server & DB Schema Setup", 3, 3),
    ("Web Client Dashboard & Twilio Alerts", 3, 4),
    ("System Integration, QA Testing & Defense", 4, 4)
]

for row_idx, (task_name, start_col, end_col) in enumerate(timeline_data, start=1):
    # Task Name Cell
    cell_task = table.cell(row_idx, 0)
    cell_task.fill.solid(); cell_task.fill.fore_color.rgb = COLOR_WHITE
    tf_t = cell_task.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = Inches(0.1)
    p_t = tf_t.paragraphs[0]
    p_t.text = task_name
    p_t.font.name = "Segoe UI"
    p_t.font.size = Pt(11)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_TEXT_MAIN
    
    # Month Columns
    for col_idx in range(1, 5):
        cell_m = table.cell(row_idx, col_idx)
        cell_m.fill.solid()
        if start_col <= col_idx <= end_col:
            # Task Active: Draw active color bar
            cell_m.fill.fore_color.rgb = COLOR_ACCENT_GREEN if row_idx % 2 == 0 else COLOR_ACCENT_BLUE
            tf_m = cell_m.text_frame
            tf_m.word_wrap = True
            p_m = tf_m.paragraphs[0]
            p_m.text = "ACTIVE DEVELOPMENT" if start_col == col_idx else ""
            p_m.font.name = "Segoe UI"
            p_m.font.size = Pt(8.5)
            p_m.font.bold = True
            p_m.font.color.rgb = COLOR_PRIMARY if row_idx % 2 == 0 else COLOR_SECONDARY
            p_m.alignment = PP_ALIGN.CENTER
        else:
            # Task Inactive: light background
            cell_m.fill.fore_color.rgb = COLOR_WHITE

add_page_number(slide13, 13)

# ==========================================
# SLIDE 14: EXPECTED OUTCOMES (4 Horizontal Outcome Bars)
# ==========================================
slide14 = prs.slides.add_slide(prs.slide_layouts[6])
apply_background(slide14)
add_slide_transition_fade(slide14)
add_slide_header(slide14, "Expected Outcomes & Project Impact")

out_w = Inches(5.6)
out_h = Inches(2.2)

# Card 1 - Yield (Green)
add_rounded_card(slide14, Inches(0.8), Inches(1.7), out_w, out_h, fill_color=COLOR_CARD_GREEN)
add_rounded_card(slide14, Inches(0.8), Inches(1.7), Inches(0.12), out_h, fill_color=COLOR_PRIMARY, border_color=None)
add_text_box(slide14, Inches(1.2), Inches(1.9), Inches(5.0), Inches(1.8), [
    {"text": "Maximizing Agricultural Yield", "font_name": "Segoe UI", "size": 15, "bold": True, "color": COLOR_PRIMARY, "space_after": 4},
    {"text": "By replacing guesswork with empirical, AI-recommended crop compatibility, the project expects a direct increase in harvest quality and density per acre, reducing crop failures.", "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MAIN}
])

# Card 2 - Soil Conservation (Blue)
add_rounded_card(slide14, Inches(6.93), Inches(1.7), out_w, out_h, fill_color=COLOR_CARD_BLUE)
add_rounded_card(slide14, Inches(6.93), Inches(1.7), Inches(0.12), out_h, fill_color=COLOR_SECONDARY, border_color=None)
add_text_box(slide14, Inches(7.33), Inches(1.9), Inches(5.0), Inches(1.8), [
    {"text": "Resource & Fertilizer Optimization", "font_name": "Segoe UI", "size": 15, "bold": True, "color": COLOR_SECONDARY, "space_after": 4},
    {"text": "By calculating the exact nutrients required for the chosen crop, farmers avoid chemical over-fertilization, preventing long-term soil acidification and reducing cost.", "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MAIN}
])

# Card 3 - Digital Telemetry Archive (Blue)
add_rounded_card(slide14, Inches(0.8), Inches(4.3), out_w, out_h, fill_color=COLOR_CARD_BLUE)
add_rounded_card(slide14, Inches(0.8), Inches(4.3), Inches(0.12), out_h, fill_color=COLOR_SECONDARY, border_color=None)
add_text_box(slide14, Inches(1.2), Inches(4.5), Inches(5.0), Inches(1.8), [
    {"text": "Empirical Regional Telemetry Data", "font_name": "Segoe UI", "size": 15, "bold": True, "color": COLOR_SECONDARY, "space_after": 4},
    {"text": "Logging telemetry continuously compiles structured geological archives of the farm, opening opportunities for academic research, soil decay studies, and predictive analytics.", "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MAIN}
])

# Card 4 - Accessibility (Green)
add_rounded_card(slide14, Inches(6.93), Inches(4.3), out_w, out_h, fill_color=COLOR_CARD_GREEN)
add_rounded_card(slide14, Inches(6.93), Inches(4.3), Inches(0.12), out_h, fill_color=COLOR_PRIMARY, border_color=None)
add_text_box(slide14, Inches(7.33), Inches(4.5), Inches(5.0), Inches(1.8), [
    {"text": "Farmer Digital Inclusivity", "font_name": "Segoe UI", "size": 15, "bold": True, "color": COLOR_PRIMARY, "space_after": 4},
    {"text": "The bilingual interface, clear color-coded dials, and automatic SMS notification system break down technology barriers, bringing modern IoT power to traditional farming.", "font_name": "Calibri", "size": 11, "color": COLOR_TEXT_MAIN}
])

add_page_number(slide14, 14)

# ==========================================
# SLIDE 15: THANK YOU SLIDE (Elegant Dark Theme Closure - Square Visual)
# ==========================================
slide15 = prs.slides.add_slide(prs.slide_layouts[6])
apply_background(slide15, COLOR_PRIMARY)
add_slide_transition_fade(slide15)

# Background Graphic/Card on the Right Side (Perfect 1:1 Aspect Ratio)
if os.path.exists(THANK_YOU_IMAGE_PATH):
    img_card = add_rounded_card(slide15, Inches(6.8), Inches(0.8), Inches(5.7), Inches(5.9), fill_color=COLOR_WHITE, border_color=None)
    # 5.5" x 5.5" image center coordinates: left = 6.8 + (5.7 - 5.5)/2 = 6.9; top = 0.8 + (5.9 - 5.5)/2 = 1.0
    slide15.shapes.add_picture(THANK_YOU_IMAGE_PATH, Inches(6.9), Inches(1.0), Inches(5.5), Inches(5.5))

# Text Card on the Left
thank_you_box = slide15.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(5.5), Inches(4.7))
tf_ty = thank_you_box.text_frame
tf_ty.word_wrap = True
tf_ty.margin_left = tf_ty.margin_right = tf_ty.margin_top = tf_ty.margin_bottom = 0

# Thank You Text
p1 = tf_ty.paragraphs[0]
p1.text = "Thank You!"
p1.font.name = "Segoe UI"
p1.font.size = Pt(52)
p1.font.bold = True
p1.font.color.rgb = COLOR_WHITE
p1.space_after = Pt(8)

p2 = tf_ty.add_paragraph()
p2.text = "Questions & Answers"
p2.font.name = "Segoe UI"
p2.font.size = Pt(21)
p2.font.bold = True
p2.font.color.rgb = COLOR_CARD_GREEN
p2.space_after = Pt(25)

p3 = tf_ty.add_paragraph()
p3.text = "IoT-Based Smart Soil Analysis and AI-Powered Crop Recommendation System"
p3.font.name = "Segoe UI"
p3.font.size = Pt(13.5)
p3.font.bold = True
p3.font.color.rgb = COLOR_WHITE
p3.space_after = Pt(15)

p4 = tf_ty.add_paragraph()
p4.text = "Students:\n▪ Jeba Takia (210104)\n▪ MD Mushfiqur Rahman Mahim (210131)"
p4.font.name = "Calibri"
p4.font.size = Pt(13)
p4.font.color.rgb = COLOR_BG
p4.space_after = Pt(15)

p5 = tf_ty.add_paragraph()
p5.text = "Department of Computer Science and Engineering\nDaffodil Institute of IT\nSession 2020-21"
p5.font.name = "Calibri"
p5.font.size = Pt(12)
p5.font.color.rgb = COLOR_BG

add_page_number(slide15, 15)

# Save presentation
prs.save("IoT_Smart_Soil_Analysis_Presentation_v2.pptx")
print("Presentation generated successfully with adjusted aspects and enhanced aesthetics!")
